# -*- coding: utf-8 -*-
"""result1—result4 导出（题面附件 3 模板；能力 P1-C3、P2-C3、P3-C4、P4-C3；约束 K13—K15）。

本文件不重算任何物理量：四张表全部取自 figures/problem_N_results.json 的 result_grid，
只做「模板对齐 + 落盘 + 回读复核」三件事：
  1. 表头（角标、半径标签、问题 4 的末列名）逐字取自 user_data/result{N}.xlsx 模板，不自拟；
  2. 数值原样写入，写前再断言一次四位小数量化（K13）；
  3. 写完重新打开每个文件逐格回读，行列数、空缺位置、数值三项全对才算通过（K14、K15）。

运行：python _utils/run_compute.py code/export_excel.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "_utils"))

from params import *  # noqa: F401,F403,E402
from utils import ensure_dirs, jsonable, save_json  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "user_data"
ELLIPSIS = {"…", "„", "...", "..", "\u22ef"}   # 模板里省略中间列的占位符，展开时丢弃


def _label(value: float):
    """整数半径写成整数（与模板的 0、2 一致），非整数保留原值。"""
    return int(round(value)) if abs(value - round(value)) < 1e-12 else round(value, 10)


def template_labels(problem: int) -> dict:
    """读模板文件，取回工作表名、角标、起始半径标签与末列标签（逐字，不自拟）。"""
    path = TEMPLATE_DIR / RESULT_FILES[problem]
    assert path.exists(), f"缺模板 {path}"
    book = openpyxl.load_workbook(path)
    out = {"sheets": list(book.sheetnames), "per_sheet": {}}
    for sheet in book.worksheets:
        head = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
        labels = [v for v in head[1:] if v not in ELLIPSIS]
        assert len(labels) >= 2, f"{path.name}/{sheet.title} 模板表头列数异常"
        out["per_sheet"][sheet.title] = {"corner": head[0], "lead": labels[:-1],
                                        "last": labels[-1], "raw": head}
    book.close()
    return out


def build_header(problem: int, sheet: str, tpl: dict, radii: list) -> tuple:
    """展开模板表头：角标 + 21 个半径标签（+ 问题 4 的移动表面列）。

    模板用「…」省略中间列，这里按 result_grid 的实际半径展开；模板给出的首列标签与
    末列标签必须与展开结果一致，否则说明输出位置与题面要求错位。
    """
    meta = tpl["per_sheet"][sheet]
    assert meta["corner"] == COL_HEADER_CORNER, \
        f"P{problem}/{sheet} 模板角标 {meta['corner']!r} 与 params 的 {COL_HEADER_CORNER!r} 不符"
    cols = [_label(v) for v in radii]
    for k, want in enumerate(meta["lead"]):
        assert cols[k] == _label(float(want)), \
            f"P{problem}/{sheet} 第 {k + 1} 个半径列 {cols[k]} 与模板 {want} 不符"
    surface = isinstance(meta["last"], str)
    if surface:                                  # 问题 4：末列是跟随动边界的药材表面
        assert cols[-1] == _label(R0_CM), f"P{problem} 数值列应止于 {R0_CM} cm"
        header = [meta["corner"], *cols, meta["last"]]
    else:                                        # 问题 1—3：末列就是最外半径
        assert cols[-1] == _label(float(meta["last"])), \
            f"P{problem}/{sheet} 末列 {cols[-1]} 与模板 {meta['last']} 不符"
        header = [meta["corner"], *cols]
    return header, surface


def collect(problem: int) -> dict:
    """从 problem_N_results.json 取出该问题要写进 xlsx 的全部单元格。"""
    src = FIG_DIR / f"problem_{problem}_results.json"
    assert src.exists(), f"缺 {src.name}，先跑 code/problem{problem}.py"
    data = json.loads(src.read_text(encoding="utf-8"))
    grid = data["result_grid"]
    radii = grid["radii_cm"]
    assert len(radii) == N_OUT_RADII, f"P{problem} 半径列数 {len(radii)} ≠ {N_OUT_RADII}"
    tpl = template_labels(problem)

    fields = {SHEET_T: "T", SHEET_C: "C"} if problem in (1, 2) else {SHEET_SINGLE: "C"}
    assert tpl["sheets"] == list(fields), \
        f"P{problem} 模板工作表 {tpl['sheets']} 与待写字段 {list(fields)} 不符"

    sheets = {}
    for sheet, key in fields.items():
        header, surface = build_header(problem, sheet, tpl, radii)
        body = grid[key]
        surf = grid["C_surface"] if surface else None
        assert (surf is not None) == surface, f"P{problem} 需要药材表面列但 JSON 未提供"
        if surface:
            assert grid["surface_column"] == header[-1] or header[-1] in grid["surface_column"], \
                f"P{problem} 表面列表头与 JSON 声明不符"
            assert len(surf) == len(body), f"P{problem} 表面列行数与表体不符"
        sheets[sheet] = {"header": header, "times_s": grid["times_s"],
                         "body": body, "surface": surf, "field": key}
    return {"problem": problem, "file": RESULT_FILES[problem], "sheets": sheets,
            "template": {k: {"sheet": k, "corner": v["corner"], "last": v["last"]}
                         for k, v in tpl["per_sheet"].items()},
            "source": src.name}


def check_quantized(spec: dict) -> dict:
    """K13：写盘前再断言一次四位小数量化，并统计空缺单元（K14 的计数依据）。"""
    stat = {}
    for sheet, s in spec["sheets"].items():
        n_val = n_blank = 0
        for row in s["body"]:
            for v in row:
                if v is None:
                    n_blank += 1
                    continue
                assert abs(v - round(v, DECIMALS)) < 1e-12, \
                    f"K13 {spec['file']}/{sheet} 出现非四位小数值 {v!r}"
                n_val += 1
        for v in (s["surface"] or []):
            assert v is not None and abs(v - round(v, DECIMALS)) < 1e-12, \
                f"K13 {spec['file']}/{sheet} 表面列出现非四位小数值 {v!r}"
            n_val += 1
        stat[sheet] = {"n_rows": len(s["body"]), "n_cols": len(s["header"]) - 1,
                       "n_values": n_val, "n_blanks": n_blank}
    return stat


def write_workbook(spec: dict) -> Path:
    """write_only 模式逐行落盘（P2 有 10800×22×2 格，避免一次性建满内存单元）。"""
    path = ROOT / spec["file"]
    book = openpyxl.Workbook(write_only=True)
    for sheet, s in spec["sheets"].items():
        ws = book.create_sheet(title=sheet)
        ws.append(s["header"])
        surf = s["surface"]
        for k, t_s in enumerate(s["times_s"]):
            row = [int(round(t_s)), *s["body"][k]]
            if surf is not None:
                row.append(surf[k])
            ws.append(row)
    book.save(path)
    book.close()
    return path


def verify_workbook(spec: dict, path: Path) -> dict:
    """回读复核：行列数、表头、时间列、每一格数值与空缺位置都要与来源一致。"""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        assert book.sheetnames == list(spec["sheets"]), \
            f"{path.name} 工作表 {book.sheetnames} 与预期 {list(spec['sheets'])} 不符"
        report = {}
        for sheet, s in spec["sheets"].items():
            ws = book[sheet]
            width = len(s["header"])
            rows = ws.iter_rows(values_only=True)
            head = list(next(rows))[:width]
            assert head == s["header"], f"{path.name}/{sheet} 表头回读不符：{head[:3]}…"
            surf = s["surface"]
            n_val = n_blank = 0
            for k, raw in enumerate(rows):
                got = list(raw)[:width]
                got += [None] * (width - len(got))     # 尾部空格可能被省略
                assert got[0] == int(round(s["times_s"][k])), \
                    f"{path.name}/{sheet} 第 {k + 1} 行时间列 {got[0]} 与来源不符"
                want = list(s["body"][k]) + ([surf[k]] if surf is not None else [])
                for j, wv in enumerate(want):
                    gv = got[j + 1]
                    if wv is None:
                        assert gv is None, \
                            f"K14 {path.name}/{sheet} 第 {k + 1} 行第 {j + 1} 列域外单元被填成 {gv!r}"
                        n_blank += 1
                    else:
                        assert gv is not None and abs(float(gv) - wv) < 5e-13, \
                            f"{path.name}/{sheet} 第 {k + 1} 行第 {j + 1} 列 {gv!r} ≠ {wv!r}"
                        n_val += 1
            assert k + 1 == len(s["times_s"]), \
                f"{path.name}/{sheet} 数据行 {k + 1} ≠ 来源 {len(s['times_s'])}"
            report[sheet] = {"n_rows": k + 1, "n_cols": width - 1,
                             "n_values": n_val, "n_blanks": n_blank,
                             "header_first3": [str(v) for v in head[:3]],
                             "header_last": str(head[-1])}
        return report
    finally:
        book.close()


def cross_check(problem: int, spec: dict, read_back: dict) -> dict:
    """把回读结果与来源 JSON 的独立声明对上：时间窗、行数、空缺总数。"""
    data = json.loads((FIG_DIR / f"problem_{problem}_results.json").read_text(encoding="utf-8"))
    grid = data["result_grid"]
    times = grid["times_s"]
    want_end = {1: P1_WINDOW_S, 2: P2_WINDOW_S}.get(problem)
    out_dt = {1: P1_OUT_DT_S, 2: P2_OUT_DT_S, 3: P3_OUT_DT_S, 4: P4_OUT_DT_S}[problem]
    assert abs(times[0] - out_dt) < 1e-9, f"P{problem} 首行时刻应为一个输出步 {out_dt} s"
    assert len(times) == int(round(times[-1] / out_dt)), f"P{problem} 行数与输出步不自洽"
    note = ""
    if want_end is not None:
        assert abs(times[-1] - want_end) < 1e-9, \
            f"P{problem} 末行 {times[-1]} s 与题面窗口 {want_end} s 不符"
        note = f"覆盖 {want_end:.0f} s 全窗，每 {out_dt:.0f} s 一行"
    else:                                  # 问题 3、4：写到跨阈值区间的上端（含 t_dry 所在步）
        t_dry = (data if problem == 3 else data["main"])["t_dry"]
        hi = t_dry["bracket_t_hi_s"]
        assert abs(times[-1] - hi) < 1e-9, \
            f"P{problem} 末行 {times[-1]} s 应为跨阈值区间上端 {hi} s"
        note = (f"写到跨阈值输出步上端 {hi:.0f} s（t_dry={t_dry['seconds']:.3f} s "
                f"= {t_dry['hours']:.4f} h 落在其内），每 {out_dt:.0f} s 一行")

    total_blank = sum(v["n_blanks"] for v in read_back.values())
    if problem == 4:
        assert total_blank == grid["n_blank_cells"], \
            f"K14 回读空缺 {total_blank} ≠ JSON 声明 {grid['n_blank_cells']}"
    else:
        assert total_blank == 0, f"K13 P{problem} 不应有空缺单元，实测 {total_blank}"
    return {"t_first_s": float(times[0]), "t_last_s": float(times[-1]),
            "out_dt_s": float(out_dt), "n_rows": len(times), "n_blanks": total_blank,
            "coverage": note}


def main() -> int:
    ensure_dirs()
    summary = {"problem": "cross", "capabilities": ["P1-C3", "P2-C3", "P3-C4", "P4-C3"],
               "constraints": ["K13", "K14", "K15"], "files": {},
               "template_dir": TEMPLATE_DIR.name,
               "note": ("数值全部取自 figures/problem_N_results.json 的 result_grid，"
                        "本文件只做模板对齐、落盘与逐格回读复核，不重算物理量。")}
    for problem in sorted(RESULT_FILES):
        spec = collect(problem)
        pre = check_quantized(spec)
        path = write_workbook(spec)
        read_back = verify_workbook(spec, path)
        for sheet, want in pre.items():
            got = read_back[sheet]
            assert all(got[k] == want[k] for k in want), \
                f"{spec['file']}/{sheet} 回读统计 {got} 与写前 {want} 不符"
        size_kb = path.stat().st_size / 1024.0
        summary["files"][spec["file"]] = {
            "problem": problem, "path": spec["file"], "source": spec["source"],
            "sheets": read_back, "size_kb": round(size_kb, 1),
            "cross_check": cross_check(problem, spec, read_back)}
        for sheet, r in read_back.items():
            print(f"[XL] {spec['file']}/{sheet}: {r['n_rows']}×{r['n_cols']}"
                  f"  数值 {r['n_values']}  空缺 {r['n_blanks']}"
                  f"  末列表头「{r['header_last']}」")
        print(f"[XL] {spec['file']}  {size_kb:.1f} KB  "
              f"{summary['files'][spec['file']]['cross_check']['coverage']}")

    n_val = sum(r["n_values"] for f in summary["files"].values() for r in f["sheets"].values())
    n_blk = sum(r["n_blanks"] for f in summary["files"].values() for r in f["sheets"].values())
    summary["totals"] = {"n_files": len(summary["files"]), "n_values": n_val,
                         "n_blanks": n_blk,
                         "size_kb": round(sum(f["size_kb"] for f in summary["files"].values()), 1)}
    print(f"[XL] 共 {summary['totals']['n_files']} 个文件、{n_val} 个数值、{n_blk} 个域外空缺"
          f"（K14），合计 {summary['totals']['size_kb']:.1f} KB；四位小数（K13）逐格回读通过")
    save_json(FIG_DIR / "export_excel_results.json", jsonable(summary))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

